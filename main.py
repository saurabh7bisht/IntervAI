import html as html_escape
import io
import json
import logging
import os
import re
import time
from pathlib import Path

from docx import Document
from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from passlib.context import CryptContext
from pypdf import PdfReader
from sqlalchemy import create_engine, text
from starlette.middleware.base import BaseHTTPMiddleware
from starlette.middleware.sessions import SessionMiddleware

try:
    from google import genai
    from google.genai import types as genai_types
except ImportError:
    genai = None
    genai_types = None

load_dotenv(Path(__file__).resolve().parent / ".env", override=False)

logger = logging.getLogger("intervai.ai")

DATABASE_URL = os.getenv(
    "DATABASE_URL", "postgresql://postgres:postgresql@localhost:5432/intervai"
).strip()
engine = create_engine(DATABASE_URL, pool_pre_ping=True)

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")


def hash_password(password: str) -> str:
    return pwd_context.hash(password)


def verify_password(password: str, password_hash: str) -> bool:
    return pwd_context.verify(password, password_hash)


app = FastAPI(title="IntervAI Enterprise", version="2.0.0")


class SecurityHeadersMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request, call_next):
        response = await call_next(request)
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["X-Frame-Options"] = "DENY"
        response.headers["X-XSS-Protection"] = "1; mode=block"
        response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
        response.headers["Permissions-Policy"] = "microphone=(self)"
        return response


app.add_middleware(SecurityHeadersMiddleware)

SESSION_SECRET = os.getenv(
    "INTERVAI_SESSION_SECRET",
    "local-development-only-change-this-secret-intervai-secure-key"
).strip()

app.add_middleware(
    SessionMiddleware,
    secret_key=SESSION_SECRET,
    same_site="lax",
    https_only=os.getenv("INTERVAI_HTTPS_ONLY", "false").lower() == "true",
    max_age=30 * 24 * 60 * 60,
)

DISPOSABLE_EMAIL_DOMAINS = {
    "tempmail.com", "mailinator.com", "10minutemail.com", "guerrillamail.com",
    "yopmail.com", "throwawaymail.com", "sharklasers.com", "trashmail.com",
    "fake.com", "test.com", "dispostable.com", "getairmail.com", "maildrop.cc",
    "inboxkitten.com", "mohmal.com", "temp-mail.org", "fakeinbox.com",
    "10mail.org", "mytempemail.com", "burnermail.io", "trashmail.net"
}

FAILED_LOGIN_ATTEMPTS = {}


def is_rate_limited(key: str, max_attempts: int = 5, window_seconds: int = 600) -> bool:
    now = time.time()
    attempts = [t for t in FAILED_LOGIN_ATTEMPTS.get(key, []) if now - t < window_seconds]
    FAILED_LOGIN_ATTEMPTS[key] = attempts
    return len(attempts) >= max_attempts


def record_failed_attempt(key: str):
    now = time.time()
    FAILED_LOGIN_ATTEMPTS.setdefault(key, []).append(now)


def clear_failed_attempts(key: str):
    FAILED_LOGIN_ATTEMPTS.pop(key, None)


GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.5-flash-lite").strip() or "gemini-3.5-flash-lite"
LAST_AI_PROVIDER = "gemini"


def get_gemini_client():
    if genai is None or not GEMINI_API_KEY:
        return None
    try:
        return genai.Client(api_key=GEMINI_API_KEY)
    except Exception as error:
        logger.warning("Gemini client initialization failed: %s", error)
        return None


def clean_json_response(raw_text):
    cleaned = (raw_text or "").strip()
    if cleaned.startswith("```"):
        lines = cleaned.splitlines()
        if lines and lines[0].startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]
        cleaned = "\n".join(lines).strip()

    if "```json" in cleaned:
        cleaned = cleaned.split("```json", 1)[1].split("```", 1)[0].strip()
    elif "```" in cleaned:
        cleaned = cleaned.split("```", 1)[1].split("```", 1)[0].strip()

    start_brace = cleaned.find("{")
    start_bracket = cleaned.find("[")
    if start_brace != -1 and (start_bracket == -1 or start_brace < start_bracket):
        end_brace = cleaned.rfind("}")
        if end_brace > start_brace:
            cleaned = cleaned[start_brace:end_brace + 1]
    elif start_bracket != -1:
        end_bracket = cleaned.rfind("]")
        if end_bracket > start_bracket:
            cleaned = cleaned[start_bracket:end_bracket + 1]

    cleaned = re.sub(r",\s*([\]}])", r"\1", cleaned)
    return cleaned


def call_ai_json_any(prompt):
    global LAST_AI_PROVIDER
    LAST_AI_PROVIDER = "gemini"

    gemini_client = get_gemini_client()
    if gemini_client is not None:
        candidate_models = list(dict.fromkeys(["gemini-3.5-flash-lite", GEMINI_MODEL, "gemini-2.5-flash", "gemini-1.5-flash"]))
        for model_name in candidate_models:
            try:
                response = gemini_client.models.generate_content(
                    model=model_name,
                    contents=prompt,
                    config=genai_types.GenerateContentConfig(
                        response_mime_type="application/json",
                        temperature=0.25,
                    ),
                )
                raw = clean_json_response(getattr(response, "text", "") or "")
                parsed = json.loads(raw)
                if isinstance(parsed, (dict, list)):
                    return parsed
            except Exception as error:
                logger.warning("Gemini JSON generation failed for model %s: %s", model_name, error)
    return None


def call_ai_json(prompt):
    result = call_ai_json_any(prompt)
    return result if isinstance(result, dict) else None


@app.get("/api/ai-status")
def ai_status():
    return {
        "gemini_sdk_installed": genai is not None,
        "gemini_configured": bool(GEMINI_API_KEY and genai is not None),
        "gemini_model": GEMINI_MODEL,
        "provider": "Google Gemini (Exclusive)",
    }


@app.post("/api/gemini/test")
async def gemini_test(request: Request):
    if genai is None:
        return {"success": False, "message": "Gemini SDK missing. Install google-genai first."}
    if not GEMINI_API_KEY:
        return {"success": False, "message": "GEMINI_API_KEY is not configured in .env or environment."}
    try:
        client = get_gemini_client()
        if client is None:
            return {"success": False, "message": "Gemini client could not initialize."}
        response = client.models.generate_content(
            model=GEMINI_MODEL,
            contents="Reply with exactly: IntervAI Gemini connection successful.",
        )
        reply = (getattr(response, "text", "") or "").strip()
        if not reply:
            return {"success": False, "message": "Gemini returned an empty response. Check model availability and API access."}
        return {"success": True, "model": GEMINI_MODEL, "reply": reply}
    except Exception as error:
        logger.warning("Gemini connectivity test failed: %s", error)
        return {"success": False, "message": "Gemini request failed. Check the API key, model name, quota, and server terminal."}


app.mount("/static", StaticFiles(directory="static"), name="static")


def create_database_tables():
    with engine.begin() as connection:
        connection.execute(text("""
            CREATE TABLE IF NOT EXISTS users (
                id SERIAL PRIMARY KEY,
                name VARCHAR(100) NOT NULL,
                email VARCHAR(255) UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """))
        connection.execute(text("""
            CREATE TABLE IF NOT EXISTS resumes (
                id SERIAL PRIMARY KEY,
                user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                filename VARCHAR(255),
                extracted_text TEXT,
                uploaded_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """))
        connection.execute(text("""
            CREATE TABLE IF NOT EXISTS job_descriptions (
                id SERIAL PRIMARY KEY,
                user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                filename VARCHAR(255),
                description_text TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """))
        connection.execute(text("""
            CREATE TABLE IF NOT EXISTS interviews (
                id SERIAL PRIMARY KEY,
                user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                role VARCHAR(150) NOT NULL,
                experience_level VARCHAR(100),
                interview_type VARCHAR(100),
                mode VARCHAR(50),
                status VARCHAR(50) DEFAULT 'created',
                started_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                completed_at TIMESTAMP
            )
        """))
        connection.execute(text("""
            CREATE TABLE IF NOT EXISTS interview_questions (
                id SERIAL PRIMARY KEY,
                interview_id INTEGER NOT NULL REFERENCES interviews(id) ON DELETE CASCADE,
                question_number INTEGER,
                question_text TEXT NOT NULL,
                question_type VARCHAR(100),
                difficulty VARCHAR(50),
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """))
        connection.execute(text("""
            CREATE TABLE IF NOT EXISTS interview_answers (
                id SERIAL PRIMARY KEY,
                question_id INTEGER NOT NULL REFERENCES interview_questions(id) ON DELETE CASCADE,
                answer_text TEXT,
                answer_duration_seconds INTEGER,
                words_per_minute FLOAT,
                evaluation_score FLOAT,
                evaluation_json TEXT,
                feedback TEXT,
                submitted_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """))
        connection.execute(text("""
            CREATE TABLE IF NOT EXISTS interview_reports (
                id SERIAL PRIMARY KEY,
                interview_id INTEGER UNIQUE NOT NULL REFERENCES interviews(id) ON DELETE CASCADE,
                overall_score FLOAT,
                technical_score FLOAT,
                communication_score FLOAT,
                accuracy_score FLOAT,
                approach_score FLOAT,
                speed_score FLOAT,
                strengths TEXT,
                weaknesses TEXT,
                suggestions TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """))
        connection.execute(text("ALTER TABLE interview_answers ADD COLUMN IF NOT EXISTS evaluation_score FLOAT"))
        connection.execute(text("ALTER TABLE interview_answers ADD COLUMN IF NOT EXISTS evaluation_json TEXT"))
        connection.execute(text("ALTER TABLE interview_answers ADD COLUMN IF NOT EXISTS feedback TEXT"))
        connection.execute(text("ALTER TABLE interviews ADD COLUMN IF NOT EXISTS question_count INTEGER DEFAULT 8"))


@app.on_event("startup")
def startup_event():
    try:
        create_database_tables()
    except Exception as error:
        logger.error("Database table creation failed: %s", error)


@app.get("/database-test")
def database_test():
    try:
        with engine.connect() as connection:
            value = connection.execute(text("SELECT 1")).scalar()
        return {"success": True, "message": "PostgreSQL connected successfully!", "database_test": value}
    except Exception as error:
        return {"success": False, "message": "PostgreSQL connection failed.", "error": str(error)}


@app.get("/tables-test")
def tables_test():
    try:
        with engine.connect() as connection:
            result = connection.execute(text("""
                SELECT table_name
                FROM information_schema.tables
                WHERE table_schema = 'public'
                ORDER BY table_name
            """))
            tables = [row[0] for row in result.fetchall()]
        return {"success": True, "tables": tables}
    except Exception as error:
        return {"success": False, "error": str(error)}


@app.post("/signup")
def signup(name: str = Form(...), email: str = Form(...), password: str = Form(...)):
    name = name.strip()
    email = email.strip().lower()

    name_parts = [p for p in name.split() if p]
    if len(name_parts) < 2:
        return {"success": False, "message": "Please enter your genuine legal Full Name (First and Last Name)."}

    if len(name) < 4 or any(c.isdigit() for c in name) or name.lower() in ("test user", "admin user", "asdf asdf", "demo user", "fake account"):
        return {"success": False, "message": "Please provide a genuine candidate name. Numeric and dummy names are not permitted."}

    if not email or "@" not in email or "." not in email:
        return {"success": False, "message": "Please enter a valid email address."}

    email_domain = email.split("@")[-1].strip().lower()
    if email_domain in DISPOSABLE_EMAIL_DOMAINS:
        return {"success": False, "message": f"Temporary and disposable emails (@{email_domain}) are not permitted. Please use your genuine personal or corporate email."}

    if len(password) < 8 or not (any(c.isalpha() for c in password) and any(c.isdigit() for c in password)):
        return {"success": False, "message": "Password must be at least 8 characters long and contain both letters and numbers."}

    try:
        with engine.connect() as connection:
            existing_user = connection.execute(
                text("SELECT id FROM users WHERE email = :email"),
                {"email": email}
            ).fetchone()

        if existing_user:
            return {"success": False, "message": "An account with this email address already exists. Please sign in."}

        password_hash = hash_password(password)

        with engine.begin() as connection:
            result = connection.execute(
                text("""
                    INSERT INTO users (name, email, password_hash)
                    VALUES (:name, :email, :password_hash)
                    RETURNING id, name, email
                """),
                {"name": name, "email": email, "password_hash": password_hash}
            )
            user = result.fetchone()

        return {
            "success": True,
            "message": "Verified candidate account created successfully!",
            "user": {"id": user.id, "name": user.name, "email": user.email}
        }
    except Exception as error:
        return {"success": False, "message": "Something went wrong while creating your account. Please try again.", "error": str(error)}


@app.post("/login")
def login(request: Request, email: str = Form(...), password: str = Form(...)):
    email = email.strip().lower()
    client_ip = request.client.host if request.client else "unknown"
    rate_key = f"{client_ip}:{email}"

    if is_rate_limited(rate_key, max_attempts=5, window_seconds=600):
        return {
            "success": False,
            "rate_limited": True,
            "message": "Too many failed login attempts detected. For your account security, sign in is temporarily locked for 10 minutes. Please try again later."
        }

    try:
        with engine.connect() as connection:
            user = connection.execute(
                text("SELECT id, name, email, password_hash FROM users WHERE email = :email"),
                {"email": email}
            ).fetchone()

        if not user or not verify_password(password, user.password_hash):
            record_failed_attempt(rate_key)
            return {"success": False, "message": "Invalid email or password."}

        clear_failed_attempts(rate_key)

        request.session["user_id"] = user.id
        request.session["user_name"] = user.name
        request.session["user_email"] = user.email

        return {
            "success": True,
            "message": "Authentication successful!",
            "user": {"id": user.id, "name": user.name, "email": user.email}
        }
    except Exception as error:
        return {"success": False, "message": "Something went wrong while logging in.", "error": str(error)}


@app.get("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse(url="/login", status_code=303)


@app.post("/api/change-password")
async def change_password(request: Request):
    user_id = request.session.get("user_id")
    if not user_id:
        return {"success": False, "message": "Please log in to change your password."}

    try:
        data = await request.json()
        current_password = str(data.get("current_password", "")).strip()
        new_password = str(data.get("new_password", "")).strip()
        confirm_password = str(data.get("confirm_password", "")).strip()

        if not current_password or not new_password:
            return {"success": False, "message": "All password fields are required."}

        if new_password != confirm_password:
            return {"success": False, "message": "New password and confirmation do not match."}

        if len(new_password) < 8 or not (any(c.isalpha() for c in new_password) and any(c.isdigit() for c in new_password)):
            return {"success": False, "message": "New password must be at least 8 characters long and contain both letters and numbers."}

        with engine.connect() as connection:
            user = connection.execute(
                text("SELECT password_hash FROM users WHERE id = :user_id"),
                {"user_id": user_id}
            ).fetchone()

        if not user or not verify_password(current_password, user.password_hash):
            return {"success": False, "message": "Current password is incorrect."}

        new_hash = hash_password(new_password)
        with engine.begin() as connection:
            connection.execute(
                text("UPDATE users SET password_hash = :hash WHERE id = :user_id"),
                {"hash": new_hash, "user_id": user_id}
            )

        return {"success": True, "message": "Password updated successfully!"}
    except Exception as error:
        return {"success": False, "message": f"Unable to update password: {str(error)}"}


@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    with open("templates/index.html", "r", encoding="utf-8") as file:
        content = file.read()

    user_id = request.session.get("user_id")
    user_name_raw = str(request.session.get("user_name", "Student")).strip()
    user_name = html_escape.escape(user_name_raw or "Student")
    user_initial = html_escape.escape((user_name_raw[0] if user_name_raw else "U").upper())

    if user_id:
        nav_html = f'''
        <div class="nav-actions">
            <a href="/interview-setup" class="nav-btn primary">+ New Interview</a>
            <a href="/dashboard" class="nav-user-pill" title="Go to Dashboard">
                <div class="user-avatar-circle-sm">{user_initial}</div>
                <span class="nav-user-name">{user_name}</span>
            </a>
        </div>
        '''
        hero_actions_html = '''
        <div class="hero-actions">
            <a href="/interview-setup" class="btn btn-primary">Start New Interview →</a>
            <a href="/dashboard" class="btn btn-secondary" style="margin-left: 10px;">Go to Dashboard</a>
        </div>
        '''
    else:
        nav_html = '''
        <div class="nav-actions">
            <a href="/login" class="nav-btn">Sign In</a>
            <a href="/signup" class="nav-btn primary">Create Account</a>
        </div>
        '''
        hero_actions_html = '''
        <div class="hero-actions">
            <a href="/signup" class="btn btn-primary">Start Preparing →</a>
            <a href="#how-it-works" class="btn btn-secondary" style="margin-left: 10px;">See How It Works</a>
        </div>
        '''

    content = content.replace("{{NAV_ACTIONS}}", nav_html)
    content = content.replace("{{HERO_ACTIONS}}", hero_actions_html)
    return content


@app.get("/interview-setup", response_class=HTMLResponse)
def interview_setup(request: Request):
    user_id = request.session.get("user_id")
    if not user_id:
        return RedirectResponse(url="/login", status_code=303)

    with open("templates/interview-setup.html", "r", encoding="utf-8") as file:
        return file.read()


@app.get("/signup", response_class=HTMLResponse)
def signup_page(request: Request):
    if request.session.get("user_id"):
        return RedirectResponse(url="/dashboard", status_code=303)

    with open("templates/signup.html", "r", encoding="utf-8") as file:
        return HTMLResponse(
            content=file.read(),
            headers={"Cache-Control": "no-cache, no-store, must-revalidate"}
        )


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    if request.session.get("user_id"):
        return RedirectResponse(url="/dashboard", status_code=303)

    with open("templates/login.html", "r", encoding="utf-8") as file:
        return HTMLResponse(
            content=file.read(),
            headers={"Cache-Control": "no-cache, no-store, must-revalidate"}
        )


@app.get("/dashboard", response_class=HTMLResponse)
def dashboard(request: Request):
    user_id = request.session.get("user_id")
    if not user_id:
        return RedirectResponse(url="/login", status_code=303)

    with open("templates/dashboard.html", "r", encoding="utf-8") as file:
        page = file.read()

    completed_interviews = 0
    answered_questions = 0
    average_score = None
    recent_rows = []
    try:
        with engine.connect() as connection:
            completed_interviews = int(connection.execute(text("""
                SELECT COUNT(*) FROM interviews
                WHERE user_id = :user_id AND status = 'completed'
            """), {"user_id": user_id}).scalar() or 0)
            answered_questions = int(connection.execute(text("""
                SELECT COUNT(*)
                FROM interview_answers a
                JOIN interview_questions q ON q.id = a.question_id
                JOIN interviews i ON i.id = q.interview_id
                WHERE i.user_id = :user_id
            """), {"user_id": user_id}).scalar() or 0)
            average_score = connection.execute(text("""
                SELECT AVG(r.overall_score)
                FROM interview_reports r
                JOIN interviews i ON i.id = r.interview_id
                WHERE i.user_id = :user_id AND i.status = 'completed'
            """), {"user_id": user_id}).scalar()
            recent_rows = connection.execute(text("""
                SELECT i.id, i.role, i.interview_type, i.mode, i.status,
                       i.started_at, i.completed_at, r.overall_score
                FROM interviews i
                LEFT JOIN interview_reports r ON r.interview_id = i.id
                WHERE i.user_id = :user_id
                ORDER BY COALESCE(i.completed_at, i.started_at) DESC, i.id DESC
                LIMIT 6
            """), {"user_id": user_id}).fetchall()
    except Exception as error:
        logger.error("Dashboard statistics query error: %s", error)

    recent_html = []
    for row in recent_rows:
        role = html_escape.escape(str(row.role or "Interview"))
        interview_type = html_escape.escape(str(row.interview_type or "Mixed"))
        mode_label = html_escape.escape(str(row.mode or "Practice").title())
        status = str(row.status or "created")
        status_label = html_escape.escape(status.replace("_", " ").title())
        date_value = row.completed_at or row.started_at
        date_label = date_value.strftime("%b %d, %Y") if date_value else "Date unavailable"
        score_label = f"{float(row.overall_score):.0f}/100" if row.overall_score is not None else "—"
        if status == "completed":
            action_link = f'<a class="history-link" href="/report/{int(row.id)}">View report →</a>'
        else:
            action_link = f'<a class="history-link" href="/interview/{int(row.id)}">Continue →</a>'
        recent_html.append(
            '<div class="history-row">'
            f'<div class="history-main"><strong>{role}</strong><span>{interview_type} · {mode_label} Mode · {html_escape.escape(date_label)}</span></div>'
            f'<div class="history-meta"><span class="history-status status-{html_escape.escape(status)}">{status_label}</span>'
            f'<strong class="history-score">{score_label}</strong>{action_link}</div></div>'
        )

    recent_content = "\n".join(recent_html) if recent_html else (
        '<div class="empty-history"><strong>No interviews yet.</strong><br>Start your first AI interview above.</div>'
    )
    score_text = f"{float(average_score):.0f}%" if average_score is not None else "—"
    user_name_raw = str(request.session.get("user_name", "Student")).strip()
    user_initial = (user_name_raw[0] if user_name_raw else "U").upper()

    replacements = {
        "{{USER_NAME}}": html_escape.escape(user_name_raw or "Student"),
        "{{USER_INITIAL}}": html_escape.escape(user_initial),
        "{{USER_EMAIL}}": html_escape.escape(str(request.session.get("user_email", ""))),
        "{{COMPLETED_INTERVIEWS}}": str(completed_interviews),
        "{{ANSWERED_QUESTIONS}}": str(answered_questions),
        "{{AVERAGE_SCORE}}": score_text,
        "{{RECENT_INTERVIEWS}}": recent_content,
    }
    for key, value in replacements.items():
        page = page.replace(key, value)
    return page


@app.get("/report/{interview_id}", response_class=HTMLResponse)
def interview_report_page(request: Request, interview_id: int):
    user_id = request.session.get("user_id")
    if not user_id:
        return RedirectResponse(url="/login", status_code=303)

    try:
        with engine.connect() as connection:
            row = connection.execute(text("""
                SELECT i.id, i.role, i.experience_level, i.interview_type,
                       i.mode, i.status, i.started_at, i.completed_at,
                       r.overall_score, r.technical_score,
                       r.communication_score, r.accuracy_score,
                       r.approach_score, r.speed_score,
                       r.strengths, r.weaknesses, r.suggestions
                FROM interviews i
                LEFT JOIN interview_reports r ON r.interview_id = i.id
                WHERE i.id = :interview_id AND i.user_id = :user_id
            """), {"interview_id": interview_id, "user_id": user_id}).fetchone()
            if not row:
                return RedirectResponse(url="/dashboard", status_code=303)
            answers = connection.execute(text("""
                SELECT q.question_number, q.question_text,
                       a.answer_text, a.evaluation_score, a.feedback,
                       a.evaluation_json
                FROM interview_questions q
                LEFT JOIN interview_answers a ON a.question_id = q.id
                WHERE q.interview_id = :interview_id
                ORDER BY q.question_number
            """), {"interview_id": interview_id}).fetchall()

        with open("templates/report.html", "r", encoding="utf-8") as file:
            page = file.read()

        def esc(value, default="Not available"):
            return html_escape.escape(str(value if value not in (None, "") else default))

        def score(value):
            return f"{float(value):.0f}" if value is not None else "—"

        is_practice = "practice" in str(row.mode or "").lower()
        date_value = row.completed_at or row.started_at
        date_label = date_value.strftime("%b %d, %Y · %I:%M %p") if date_value else "Date unavailable"
        
        answer_cards = []
        for answer in answers:
            try:
                evaluation = json.loads(answer.evaluation_json or "{}")
            except (TypeError, ValueError, json.JSONDecodeError):
                evaluation = {}
            strengths = evaluation.get("strengths", [])
            improvements = evaluation.get("improvements", [])
            missing_concepts = evaluation.get("missing_concepts", [])
            what_to_answer = evaluation.get("what_to_answer", "")
            recommended_method = evaluation.get("recommended_method", "")
            method_guide = evaluation.get("method_guide", "")
            input_mode = evaluation.get("input_mode", "typed")

            voice_badge = '<span class="badge badge-voice">🎙 SPOKEN</span>' if input_mode == "voice" else ""
            score_badge = f'<span class="badge badge-score">Accuracy: {score(answer.evaluation_score)}/100</span>'

            strengths_html = "".join(f"<li>{esc(item)}</li>" for item in strengths[:5]) if isinstance(strengths, list) else ""
            improvements_html = "".join(f"<li>{esc(item)}</li>" for item in improvements[:5]) if isinstance(improvements, list) else ""
            if missing_concepts and isinstance(missing_concepts, list):
                improvements_html += "".join(f"<li>{esc(item)}</li>" for item in missing_concepts[:5])

            detail_cols = ""
            if strengths_html or improvements_html:
                detail_cols = (
                    '<div class="detail-cols" style="margin-top:16px;">'
                    f'<div class="detail-col str"><h4>Strengths</h4><ul>{strengths_html}</ul></div>'
                    f'<div class="detail-col imp"><h4>Areas to Improve</h4><ul>{improvements_html}</ul></div>'
                    '</div>'
                )

            coaching_highlight = ""
            if is_practice:
                method_box = ""
                if recommended_method:
                    method_box = (
                        f'<div style="margin-bottom:12px;padding:12px;border-radius:10px;background:rgba(79,126,248,0.08);border:1px solid rgba(79,126,248,0.2);">'
                        f'<div style="font-size:12px;font-weight:800;color:#93c5fd;text-transform:uppercase;letter-spacing:0.05em;margin-bottom:4px;">🧠 Recommended Method & Framework: <span style="color:#ffffff;">{esc(recommended_method)}</span></div>'
                        f'<div style="font-size:13px;color:#cbd5e1;line-height:1.5;">{esc(method_guide)}</div>'
                        f'</div>'
                    )
                what_box = ""
                if what_to_answer:
                    what_box = (
                        f'<div style="margin-bottom:12px;padding:12px;border-radius:10px;background:rgba(16,185,129,0.06);border:1px solid rgba(16,185,129,0.2);">'
                        f'<div style="font-size:12px;font-weight:800;color:#34d399;text-transform:uppercase;letter-spacing:0.05em;margin-bottom:4px;">🌟 What You Should Answer (Ideal Key Points):</div>'
                        f'<div style="font-size:13px;color:#e2e8f0;line-height:1.6;white-space:pre-wrap;">{esc(what_to_answer)}</div>'
                        f'</div>'
                    )
                coaching_highlight = f'<div class="practice-dossier-box" style="margin-top:16px;">{method_box}{what_box}</div>'

            answer_cards.append(
                '<article class="answer-card">'
                f'<div class="answer-card-header">'
                f'<span class="q-label">{"PRACTICE QUESTION" if is_practice else "QUESTION"} {int(answer.question_number or 0)}</span>'
                f'<div class="q-badges">{score_badge}{voice_badge}</div>'
                f'</div>'
                '<div class="answer-card-body">'
                f'<h3>{esc(answer.question_text)}</h3>'
                '<div class="answer-section-label">Your Submitted Answer</div>'
                f'<div class="answer-text">{esc(answer.answer_text, "No answer recorded")}</div>'
                f'{coaching_highlight}'
                f'<div class="feedback-block"><div style="font-size:11px;font-weight:800;text-transform:uppercase;color:#94a3b8;letter-spacing:0.06em;margin-bottom:4px;">Interviewer Evaluation</div><p>{esc(answer.feedback, "Detailed feedback unavailable")}</p></div>'
                f'{detail_cols}'
                '</div></article>'
            )
        answer_content = "\n".join(answer_cards) if answer_cards else '<div style="padding:24px;color:#94a3b8;text-align:center;">No answers have been recorded for this interview session.</div>'

        report_title = "Practice Coaching <span class='gradient-text'>Dossier</span>" if is_practice else "Interview <span class='gradient-text'>Report</span>"
        report_badge = "✦ Practice Coaching History" if is_practice else "✦ Interview Performance Report"
        body_class = "practice-mode" if is_practice else "real-mode"

        values = {
            "{{USER_NAME}}": esc(request.session.get("user_name", "Student")),
            "{{INTERVIEW_ID}}": str(row.id),
            "{{ROLE}}": esc(row.role),
            "{{LEVEL}}": esc(row.experience_level),
            "{{INTERVIEW_TYPE}}": esc(row.interview_type),
            "{{MODE}}": esc(row.mode),
            "{{DATE}}": esc(date_label),
            "{{STATUS}}": esc(str(row.status or "created").replace("_", " ").title()),
            "{{OVERALL_SCORE}}": score(row.overall_score),
            "{{TECHNICAL_SCORE}}": score(row.technical_score),
            "{{COMMUNICATION_SCORE}}": score(row.communication_score),
            "{{ACCURACY_SCORE}}": score(row.accuracy_score),
            "{{APPROACH_SCORE}}": score(row.approach_score),
            "{{SPEED_SCORE}}": score(row.speed_score),
            "{{STRENGTHS}}": esc(row.strengths, "Practice regularly to improve strengths."),
            "{{WEAKNESSES}}": esc(row.weaknesses, "Review question-level feedback."),
            "{{SUGGESTIONS}}": esc(row.suggestions, "Study the recommended frameworks and model answers above."),
            "{{ANSWER_DETAILS}}": answer_content,
            "{{REPORT_TITLE}}": report_title,
            "{{REPORT_BADGE}}": report_badge,
            "{{BODY_CLASS}}": body_class,
        }
        for key, value in values.items():
            page = page.replace(key, value)
        return page
    except Exception as error:
        logger.error("Interview report page error: %s", error)
        return RedirectResponse(url="/dashboard", status_code=303)


def extract_pdf_text(file_bytes):
    pdf_file = io.BytesIO(file_bytes)
    reader = PdfReader(pdf_file)
    text_content = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_content += page_text + "\n"
    return text_content.strip()


def extract_docx_text(file_bytes):
    doc_file = io.BytesIO(file_bytes)
    document = Document(doc_file)
    text_content = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_content += paragraph.text + "\n"
    return text_content.strip()


async def extract_file_text(uploaded_file):
    if uploaded_file is None:
        return ""
    filename = (uploaded_file.filename or "").lower()
    allowed_extensions = (".pdf", ".doc", ".docx")
    if not filename.endswith(allowed_extensions):
        raise ValueError("Only PDF, DOC and DOCX files are supported.")

    file_bytes = await uploaded_file.read()
    if not file_bytes:
        raise ValueError("The uploaded file is empty.")

    if len(file_bytes) > 10 * 1024 * 1024:
        raise ValueError("File size exceeds 10MB limit. Please upload a smaller document.")

    if filename.endswith(".pdf"):
        if not file_bytes.startswith(b"%PDF"):
            raise ValueError("Invalid PDF format: file is not a valid PDF document.")
        return extract_pdf_text(file_bytes)

    if filename.endswith(".docx"):
        if not file_bytes.startswith(b"PK"):
            raise ValueError("Invalid Word format: file is not a valid DOCX document.")
        return extract_docx_text(file_bytes)

    if filename.endswith(".doc"):
        raise ValueError("Legacy .doc files are not supported. Please convert the file to .docx or PDF.")

    return ""


def generate_ai_questions(
    role, experience_level, interview_type, mode, resume_text, job_description_text, question_count=3
):
    resume_context = (resume_text or "")[:14000]
    jd_context = (job_description_text or "")[:14000]

    type_focus = {
        "Technical": "Ask industry-tested technical questions: core framework mechanics, coding logic, debugging real problems, database/API choices, and architecture trade-offs.",
        "HR": "Ask situational, culture fit, compensation, motivation, conflict handling, and behavioral questions used by top company recruiters.",
        "Behavioral": "Ask STAR-focused behavioral challenges: high-stakes deadlines, managing conflict, overcoming failure, and ownership.",
        "System Design": "Ask realistic system design problems: scaling, data modeling, caching, latency vs throughput trade-offs, and handling system failure.",
        "Case Study": "Present real-world business & technical problems. Test structured reasoning, trade-offs, and metric-driven problem solving.",
        "Project": "Deep-dive into the candidate's actual projects from their resume. Challenge their technical decisions, architecture, and edge cases.",
        "Mixed": "Mix high-impact technical problems, behavioral scenarios, and role-specific challenges as done in top tech company interview rounds.",
    }.get(interview_type, "Ask realistic industry questions for this role.")

    prompt = f"""You are an elite Lead Interviewer & Technical Hiring Manager at a top-tier company conducting a live interview.
Your task is to review the candidate's profile and the job description, and generate opening industry-standard, real-world interview questions.

CANDIDATE TARGET ROLE: {role}
EXPERIENCE LEVEL: {experience_level}
INTERVIEW TRACK: {interview_type}
INTERVIEW MODE: {mode}

TARGET JOB DESCRIPTION:
{jd_context or "Not provided"}

CANDIDATE RESUME & BACKGROUND:
{resume_context or "Not provided"}

INTERVIEW FOCUS:
{type_focus}

CRITICAL INTERVIEW GUIDELINES:
1. Act as a real human interviewer at a top tech company. Do NOT just ask shallow trivia.
2. Formulate authentic, industry-standard, and company-tested questions that real companies ask for this role and seniority.
3. Incorporate the candidate's background from their resume AND the company's JD requirements into practical scenarios (e.g. system design choices, code debugging, trade-offs, real incidents).
4. Generate the opening 3 structured questions:
   - Question 1: Role-specific opening deep dive into a core project or domain skill from their background.
   - Question 2: Practical, real-world industry technical problem or scenario directly required by the JD.
   - Question 3: Core architectural, algorithmic, or situational problem testing depth and decision-making.

Return ONLY a valid JSON array of objects with exactly these keys:
[
  {{
    "question_number": 1,
    "question_text": "...",
    "question_type": "{interview_type}",
    "difficulty": "Foundation"
  }},
  {{
    "question_number": 2,
    "question_text": "...",
    "question_type": "{interview_type}",
    "difficulty": "Intermediate"
  }},
  {{
    "question_number": 3,
    "question_text": "...",
    "question_type": "{interview_type}",
    "difficulty": "Advanced"
  }}
]
"""
    try:
        questions = call_ai_json_any(prompt)
        if not isinstance(questions, list):
            return None

        valid_questions = []
        for index, item in enumerate(questions[:4]):
            if not isinstance(item, dict):
                continue
            question_text = str(item.get("question_text", "")).strip()
            if not question_text:
                continue
            valid_questions.append({
                "question_number": index + 1,
                "question_text": question_text,
                "question_type": str(item.get("question_type", interview_type)).strip() or interview_type,
                "difficulty": str(item.get("difficulty", "Intermediate")).strip() or "Intermediate"
            })

        if len(valid_questions) < 2:
            return None

        return valid_questions
    except Exception as error:
        logger.error("AI question generation error: %s", error)
        return None


def save_questions_for_interview(interview_id, questions):
    with engine.begin() as connection:
        for question in questions:
            connection.execute(
                text("""
                    INSERT INTO interview_questions (interview_id, question_number, question_text, question_type, difficulty)
                    VALUES (:interview_id, :question_number, :question_text, :question_type, :difficulty)
                """),
                {
                    "interview_id": interview_id,
                    "question_number": question["question_number"],
                    "question_text": question["question_text"],
                    "question_type": question["question_type"],
                    "difficulty": question["difficulty"],
                }
            )


def get_local_questions(role, experience_level, interview_type, question_count=4):
    role_name = role or "your target role"
    question_bank = {
        "Technical": [
            f"Explain how you design, test, and deploy a production-grade service as a {role_name}.",
            f"Describe a complex technical bug or system failure you resolved in {role_name} work. How did you diagnose it?",
            f"How do you approach database performance, caching, and query optimization in high-throughput applications?",
            "What security practices and vulnerability mitigations do you enforce in production systems?"
        ],
        "HR": [
            f"Walk me through your background and why you are uniquely qualified for the {role_name} role.",
            "Tell me about a high-stakes project with tight deadlines. How did you prioritize competing demands?",
            "Describe a time you received constructive criticism from leadership. How did you adapt?",
            "What values and culture enable you to perform at your peak in an engineering team?"
        ],
        "Behavioral": [
            "Tell me about a technical disagreement with a colleague or lead. How did you navigate it?",
            "Describe a situation where a project failed or went off-track. What did you take ownership of?",
            "Walk me through an instance where you mentored a junior engineer or drove technical consensus.",
            "Describe how you adapted when key project requirements changed days before a release."
        ],
        "Project": [
            f"Walk me through an end-to-end project that exemplifies your expertise as a {role_name}.",
            "What critical architectural trade-offs did you evaluate, and why did you choose your final approach?",
            "How did you monitor performance, error rates, and user impact after launch?",
            "If you had to re-architect this system today from scratch, what would you design differently?"
        ],
        "System Design": [
            f"Design a scalable, highly available distributed system relevant to {role_name}.",
            "How would you identify bottlenecks, partition data, and handle sudden 10x traffic spikes?",
            "How would you guarantee data consistency, fault tolerance, and disaster recovery across regions?",
            "What logging, tracing, and metric alerts would you establish for 99.99% availability?"
        ],
        "Case Study": [
            f"Suppose a key service for a {role_name} application starts throwing 500 errors during peak hours. How do you triage?",
            "How do you evaluate whether to build an internal solution versus adopting a managed vendor service?",
            "How do you define and communicate technical debt versus feature velocity to product managers?",
            "How do you measure success and user adoption for newly launched technical capabilities?"
        ],
        "Mixed": [
            f"Tell me about your core technical domain and what makes you effective as a {role_name}.",
            f"Walk me through the toughest architectural challenge you have faced and how you solved it.",
            "How do you maintain code quality, automated test coverage, and continuous delivery in fast-paced teams?",
            "Why is this specific company and role the right next step for your professional trajectory?"
        ]
    }

    questions = question_bank.get(interview_type, question_bank["Mixed"])
    selected_questions = questions[:question_count]
    return [
        {
            "question_number": index + 1,
            "question_text": question,
            "question_type": interview_type,
            "difficulty": "Foundation" if index == 0 else ("Intermediate" if index == 1 else "Advanced")
        }
        for index, question in enumerate(selected_questions)
    ]


def ensure_interview_questions(interview_id, role, experience_level, interview_type, question_count=3):
    with engine.begin() as connection:
        existing = connection.execute(
            text("SELECT id FROM interview_questions WHERE interview_id = :interview_id ORDER BY question_number"),
            {"interview_id": interview_id}
        ).fetchall()
        if existing:
            return

        questions = get_local_questions(role, experience_level, interview_type, question_count)
        for question in questions:
            connection.execute(
                text("""
                    INSERT INTO interview_questions (interview_id, question_number, question_text, question_type, difficulty)
                    VALUES (:interview_id, :question_number, :question_text, :question_type, :difficulty)
                """),
                {
                    "interview_id": interview_id,
                    "question_number": question["question_number"],
                    "question_text": question["question_text"],
                    "question_type": question["question_type"],
                    "difficulty": question["difficulty"]
                }
            )


def generate_questions_with_fallback(
    interview_id, role, experience_level, interview_type, mode, resume_text, job_description_text, question_count=3
):
    ai_questions = generate_ai_questions(
        role, experience_level, interview_type, mode, resume_text, job_description_text, question_count
    )
    if ai_questions:
        save_questions_for_interview(interview_id, ai_questions)
        return "ai"

    ensure_interview_questions(interview_id, role, experience_level, interview_type, question_count)
    return "local"


def evaluate_interview_answer(
    question_text, answer_text, role, level, interview_type, mode, resume_text="", jd_text="", previous_qa=None
):
    history_str = ""
    if previous_qa:
        history_str = "\n".join(
            f"- Q{qa.get('num', idx+1)}: {qa.get('q', '')} -> A: {str(qa.get('a', ''))[:200]}"
            for idx, qa in enumerate(previous_qa)
        )

    resume_context = (resume_text or "")[:4000]
    jd_context = (jd_text or "")[:4000]

    prompt = f"""
You are an elite Senior Hiring Manager & Principal Interviewer at a top tech company evaluating a candidate's live interview response.
Evaluate the answer with realistic industry rigor, accuracy, and constructive coaching.

TARGET ROLE: {role}
SENIORITY LEVEL: {level}
INTERVIEW TRACK: {interview_type}
INTERVIEW MODE: {mode}

JOB DESCRIPTION:
{jd_context or "Not provided"}

CANDIDATE RESUME CONTEXT:
{resume_context or "Not provided"}

SESSION CONTEXT:
{history_str or "First question."}

QUESTION ASKED:
{question_text}

CANDIDATE ANSWER:
{answer_text}

TASK:
Evaluate the answer accurately against industry standards.
Determine what the candidate should/could have answered, the best method/framework to approach this question, and whether you as the interviewer want to probe deeper (interviewer interest).

Return ONLY valid JSON with exactly these keys:
- score: integer 0-100 (overall score for this response)
- technical_accuracy: integer 0-100 (correctness of claims and concepts)
- clarity: integer 0-100 (structure, articulation, communication)
- problem_solving: integer 0-100 (engineering judgment, approach, reasoning)
- feedback: string (concise direct critique of candidate's answer)
- what_to_answer: string (2-4 concrete key points and solutions an ideal candidate at a top company would cover for this question)
- recommended_method: string (e.g., 'STAR Method (Situation, Task, Action, Result)', 'Top-Down Architecture Decomposition', 'Problem-Context-Solution', 'Trade-Off & Complexity Analysis')
- method_guide: string (clear 1-2 sentence actionable guide on how to structure answers using this method)
- strengths: array of short strings (specific strengths observed)
- improvements: array of short strings (specific missing points or gaps)
- missing_concepts: array of short strings (crucial concepts/keywords omitted)
- interviewer_interest: string ('high' if candidate gave an interesting, deep, or promising answer worth probing further; 'moderate' if standard; 'low' if off-track)
- recommend_follow_up: boolean (true if candidate made an interesting claim, proposed a specific design, or had an intriguing gap that a real interviewer would follow up on)
- follow_up_focus: string (what specific angle to probe next if follow-up is recommended)
- next_action: string ('follow_up' to probe depth, 'new_core_topic' to test another JD skill, or 'conclude' if sufficient evidence is gathered)
- suggested_difficulty: string ('Foundation', 'Intermediate', or 'Advanced')
"""
    result = call_ai_json(prompt)
    if result and isinstance(result, dict):
        try:
            score_val = max(0, min(100, int(result.get("score", 60))))
            result["score"] = score_val
            for key in ("technical_accuracy", "clarity", "problem_solving"):
                result[key] = max(0, min(100, int(result.get(key, score_val))))
            result["relevance_completeness"] = result.get("technical_accuracy", score_val)
            result["feedback"] = str(result.get("feedback", "Answer evaluated."))
            result["what_to_answer"] = str(result.get("what_to_answer", "Provide clear technical definitions, structured reasoning, and concrete trade-offs."))
            result["recommended_method"] = str(result.get("recommended_method", "STAR / Structured Problem Breakdown"))
            result["method_guide"] = str(result.get("method_guide", "State the core problem clearly, outline the approach, and justify your engineering decisions."))
            result["missing_concepts"] = [str(x) for x in result.get("missing_concepts", []) if x] if isinstance(result.get("missing_concepts"), list) else []
            result["strengths"] = [str(x) for x in result.get("strengths", []) if x] if isinstance(result.get("strengths"), list) else []
            result["improvements"] = [str(x) for x in result.get("improvements", result.get("weaknesses", [])) if x] if isinstance(result.get("improvements", result.get("weaknesses", [])), list) else []
            result["interviewer_interest"] = str(result.get("interviewer_interest", "moderate")).strip().lower()
            result["recommend_follow_up"] = bool(result.get("recommend_follow_up", False))
            result["follow_up_focus"] = str(result.get("follow_up_focus", ""))
            result["next_action"] = str(result.get("next_action", "new_core_topic")).strip().lower()
            result["suggested_difficulty"] = str(result.get("suggested_difficulty", "Intermediate")).strip()
            result["ai_evaluated"] = True
            result["ai_provider"] = "gemini"
            return result
        except (TypeError, ValueError) as err:
            logger.warning("Error parsing Gemini evaluation JSON: %s", err)

    word_count = len(answer_text.split())
    score = 35 if word_count < 8 else (55 if word_count < 25 else 70)
    return {
        "score": score,
        "technical_accuracy": score,
        "relevance_completeness": score,
        "clarity": score,
        "problem_solving": score,
        "feedback": "Answer recorded and evaluated.",
        "what_to_answer": "Structure your answer with clear definitions, practical code/system examples, and mention edge cases and trade-offs.",
        "recommended_method": "STAR Method (Situation, Task, Action, Result)",
        "method_guide": "State the situation, clarify your task, explain specific actions taken, and highlight measurable results.",
        "missing_concepts": [],
        "strengths": ["Answer was submitted and recorded."],
        "improvements": ["Elaborate on specific architectural details, edge cases, and metrics."],
        "interviewer_interest": "moderate",
        "recommend_follow_up": False,
        "follow_up_focus": "",
        "next_action": "new_core_topic",
        "suggested_difficulty": "Intermediate",
        "ai_evaluated": False,
        "ai_provider": "gemini"
    }


def generate_adaptive_question(
    role, level, interview_type, mode, question_text, answer_text, evaluation, question_number,
    resume_text="", jd_text="", previous_qa=None
):
    history_str = ""
    if previous_qa:
        history_str = "\n".join(
            f"- Q{qa.get('num', idx+1)}: {qa.get('q', '')} -> A: {str(qa.get('a', ''))[:250]}"
            for idx, qa in enumerate(previous_qa)
        )

    resume_context = (resume_text or "")[:4000]
    jd_context = (jd_text or "")[:4000]

    prompt = f"""
You are an elite Technical Hiring Manager conducting a live interview at a top company.
You just listened to the candidate's response. Now decide on the next question (#{question_number}).

ROLE: {role}
LEVEL: {level}
TRACK: {interview_type}
MODE: {mode}

LATEST QUESTION:
{question_text}

CANDIDATE'S LATEST ANSWER:
{answer_text}

YOUR EVALUATION OF THIS ANSWER:
{json.dumps(evaluation, ensure_ascii=False)}

FOLLOW-UP FOCUS (IF ANY):
{evaluation.get('follow_up_focus', '')}

INTERVIEW HISTORY:
{history_str or "First question completed."}

JOB DESCRIPTION REQUIREMENTS:
{jd_context or "Not provided"}

CANDIDATE RESUME:
{resume_context or "Not provided"}

INTERVIEWER INSTRUCTIONS:
1. If the candidate gave an interesting, deep, or promising answer, or claimed experience with a specific technology/project, BE CURIOUS: ask an insightful, challenging follow-up question probing edge cases, failure scenarios, or architectural trade-offs.
2. If the candidate answered adequately or follow-up is not needed: smoothly pivot to another critical skill, industry-standard problem, or scenario required by the JD.
3. Keep the tone realistic, professional, and conversational—just like a real interviewer.
4. Do NOT repeat questions or topics already asked.

Return ONLY valid JSON with:
{{"question_text": "...", "question_type": "{interview_type}", "difficulty": "{evaluation.get('suggested_difficulty', 'Intermediate')}"}}
"""
    result = call_ai_json(prompt)
    if not result or not isinstance(result, dict):
        return None
    question_text_next = str(result.get("question_text", "")).strip()
    if not question_text_next:
        return None
    return {
        "question_number": int(question_number),
        "question_text": question_text_next,
        "question_type": str(result.get("question_type", interview_type)).strip() or interview_type,
        "difficulty": str(result.get("difficulty", evaluation.get("suggested_difficulty", "Intermediate"))).strip() or "Intermediate"
    }


def calculate_wpm(answer_text, duration_seconds):
    words = len(answer_text.split())
    if not duration_seconds or duration_seconds <= 0:
        return 0.0
    minutes = duration_seconds / 60
    if minutes <= 0:
        return 0.0
    return round(words / minutes, 2)


@app.get("/interview/{interview_id}", response_class=HTMLResponse)
def live_interview_page(request: Request, interview_id: int):
    user_id = request.session.get("user_id")
    if not user_id:
        return RedirectResponse(url="/login", status_code=303)

    try:
        with engine.connect() as connection:
            interview = connection.execute(
                text("""
                    SELECT id, role, experience_level, interview_type, mode, status
                    FROM interviews
                    WHERE id = :interview_id AND user_id = :user_id
                """),
                {"interview_id": interview_id, "user_id": user_id}
            ).fetchone()

        if not interview:
            return RedirectResponse(url="/dashboard", status_code=303)

        ensure_interview_questions(
            interview.id, interview.role, interview.experience_level, interview.interview_type
        )

        with engine.begin() as connection:
            connection.execute(
                text("""
                    UPDATE interviews
                    SET status = 'in_progress',
                        started_at = COALESCE(started_at, CURRENT_TIMESTAMP)
                    WHERE id = :interview_id AND user_id = :user_id AND status != 'completed'
                """),
                {"interview_id": interview_id, "user_id": user_id}
            )

        with open("templates/interview.html", "r", encoding="utf-8") as file:
            html = file.read()

        return (
            html
            .replace("{{INTERVIEW_ID}}", str(interview.id))
            .replace("{{ROLE}}", str(interview.role))
            .replace("{{LEVEL}}", str(interview.experience_level or ""))
            .replace("{{INTERVIEW_TYPE}}", str(interview.interview_type or ""))
            .replace("{{MODE}}", str(interview.mode or ""))
        )
    except Exception as error:
        logger.error("Live interview page error: %s", error)
        return RedirectResponse(url="/dashboard", status_code=303)


@app.get("/api/interview/{interview_id}/current-question")
def current_interview_question(request: Request, interview_id: int):
    user_id = request.session.get("user_id")
    if not user_id:
        return {"success": False, "message": "Please login first.", "redirect": "/login"}

    try:
        with engine.connect() as connection:
            interview = connection.execute(
                text("""
                    SELECT id, role, experience_level, interview_type, mode, status
                    FROM interviews
                    WHERE id = :interview_id AND user_id = :user_id
                """),
                {"interview_id": interview_id, "user_id": user_id}
            ).fetchone()

            if not interview:
                return {"success": False, "message": "Interview not found."}

            question = connection.execute(
                text("""
                    SELECT q.id, q.question_number, q.question_text, q.question_type, q.difficulty
                    FROM interview_questions q
                    WHERE q.interview_id = :interview_id
                      AND NOT EXISTS (
                          SELECT 1 FROM interview_answers a WHERE a.question_id = q.id
                      )
                    ORDER BY q.question_number
                    LIMIT 1
                """),
                {"interview_id": interview_id}
            ).fetchone()

            answered_count = connection.execute(
                text("""
                    SELECT COUNT(*)
                    FROM interview_answers a
                    JOIN interview_questions q ON q.id = a.question_id
                    WHERE q.interview_id = :interview_id
                """),
                {"interview_id": interview_id}
            ).scalar() or 0

        if not question:
            return {
                "success": True,
                "finished": True,
                "answered_count": int(answered_count),
                "message": "All interview questions have been completed."
            }

        return {
            "success": True,
            "finished": False,
            "answered_count": int(answered_count),
            "question": {
                "id": question.id,
                "number": question.question_number,
                "question_number": question.question_number,
                "text": question.question_text,
                "question_text": question.question_text,
                "type": question.question_type,
                "difficulty": question.difficulty
            },
            "interview": {
                "role": interview.role,
                "level": interview.experience_level,
                "type": interview.interview_type,
                "mode": interview.mode
            }
        }
    except Exception as error:
        logger.error("Current question error: %s", error)
        return {"success": False, "message": "Unable to load the interview question."}


@app.post("/api/interview/{interview_id}/answer")
async def submit_interview_answer(request: Request, interview_id: int):
    user_id = request.session.get("user_id")
    if not user_id:
        return {"success": False, "message": "Please login first.", "redirect": "/login"}

    try:
        data = await request.json()
        question_id = int(data.get("question_id", 0))
        answer_text = str(data.get("answer_text", "")).strip()
        duration_seconds = int(data.get("duration_seconds", 0) or 0)

        if not question_id:
            return {"success": False, "message": "Question ID is missing."}
        if not answer_text:
            return {"success": False, "message": "Please provide an answer."}

        with engine.begin() as connection:
            valid_question = connection.execute(
                text("""
                    SELECT q.id, q.question_number, q.question_text,
                           i.role, i.experience_level, i.interview_type, i.mode,
                           COALESCE(i.question_count, 8) AS question_count
                    FROM interview_questions q
                    JOIN interviews i ON i.id = q.interview_id
                    WHERE q.id = :question_id AND q.interview_id = :interview_id AND i.user_id = :user_id
                """),
                {"question_id": question_id, "interview_id": interview_id, "user_id": user_id}
            ).fetchone()

            if not valid_question:
                return {"success": False, "message": "Invalid interview question."}

            already_answered = connection.execute(
                text("SELECT id FROM interview_answers WHERE question_id = :question_id"),
                {"question_id": question_id}
            ).fetchone()

            if already_answered:
                return {"success": False, "message": "This question has already been answered."}

            resume_row = connection.execute(
                text("SELECT extracted_text FROM resumes WHERE user_id = :user_id ORDER BY id DESC LIMIT 1"),
                {"user_id": user_id}
            ).fetchone()
            resume_text = resume_row.extracted_text if resume_row else ""

            jd_row = connection.execute(
                text("SELECT description_text FROM job_descriptions WHERE user_id = :user_id ORDER BY id DESC LIMIT 1"),
                {"user_id": user_id}
            ).fetchone()
            jd_text = jd_row.description_text if jd_row else ""

            past_qa_rows = connection.execute(
                text("""
                    SELECT q.question_number, q.question_text, a.answer_text
                    FROM interview_questions q
                    JOIN interview_answers a ON a.question_id = q.id
                    WHERE q.interview_id = :interview_id
                    ORDER BY q.question_number
                """),
                {"interview_id": interview_id}
            ).fetchall()
            previous_qa = [
                {"num": row.question_number, "q": row.question_text, "a": row.answer_text}
                for row in past_qa_rows
            ]

            input_mode = str(data.get("input_mode", "typed")).strip().lower()
            if input_mode not in ("voice", "typed"):
                input_mode = "typed"

            wpm = calculate_wpm(answer_text, duration_seconds)
            evaluation = evaluate_interview_answer(
                valid_question.question_text,
                answer_text,
                valid_question.role,
                valid_question.experience_level or "Unspecified",
                valid_question.interview_type or "Mixed",
                valid_question.mode or "Practice",
                resume_text=resume_text,
                jd_text=jd_text,
                previous_qa=previous_qa
            )
            evaluation["input_mode"] = input_mode

            answer_result = connection.execute(
                text("""
                    INSERT INTO interview_answers (
                        question_id, answer_text, answer_duration_seconds,
                        words_per_minute, evaluation_score, evaluation_json, feedback
                    )
                    VALUES (
                        :question_id, :answer_text, :duration_seconds,
                        :wpm, :evaluation_score, :evaluation_json, :feedback
                    )
                    RETURNING id
                """),
                {
                    "question_id": question_id,
                    "answer_text": answer_text,
                    "duration_seconds": duration_seconds,
                    "wpm": wpm,
                    "evaluation_score": evaluation.get("score", 0),
                    "evaluation_json": json.dumps(evaluation, ensure_ascii=False),
                    "feedback": evaluation.get("feedback", "")
                }
            )
            answer_row = answer_result.fetchone()

            answered_count = connection.execute(
                text("""
                    SELECT COUNT(*)
                    FROM interview_answers a
                    JOIN interview_questions q ON q.id = a.question_id
                    WHERE q.interview_id = :interview_id
                """),
                {"interview_id": interview_id}
            ).scalar() or 0

            interviewer_interest = evaluation.get("interviewer_interest", "moderate")
            recommend_follow_up = evaluation.get("recommend_follow_up", False)
            next_action = evaluation.get("next_action", "new_core_topic")

            adaptive_added = False
            interview_finished = False

            if int(answered_count) >= 8 or (
                int(answered_count) >= 5
                and next_action == "conclude"
                and not recommend_follow_up
                and interviewer_interest != "high"
            ):
                connection.execute(
                    text("""
                        DELETE FROM interview_questions
                        WHERE interview_id = :interview_id
                          AND NOT EXISTS (
                              SELECT 1 FROM interview_answers a WHERE a.question_id = interview_questions.id
                          )
                    """),
                    {"interview_id": interview_id}
                )
                interview_finished = True
            else:
                next_q_num = int(answered_count) + 1
                current_qa_list = previous_qa + [{
                    "num": valid_question.question_number,
                    "q": valid_question.question_text,
                    "a": answer_text
                }]
                adaptive_question = generate_adaptive_question(
                    valid_question.role,
                    valid_question.experience_level or "Unspecified",
                    valid_question.interview_type or "Mixed",
                    valid_question.mode or "Practice",
                    valid_question.question_text,
                    answer_text,
                    evaluation,
                    next_q_num,
                    resume_text=resume_text,
                    jd_text=jd_text,
                    previous_qa=current_qa_list
                )
                if adaptive_question:
                    next_unanswered = connection.execute(
                        text("""
                            SELECT id FROM interview_questions
                            WHERE interview_id = :interview_id
                              AND NOT EXISTS (
                                  SELECT 1 FROM interview_answers a WHERE a.question_id = interview_questions.id
                              )
                            ORDER BY question_number
                            LIMIT 1
                        """),
                        {"interview_id": interview_id}
                    ).fetchone()

                    if next_unanswered:
                        connection.execute(
                            text("""
                                UPDATE interview_questions
                                SET question_text = :question_text,
                                    question_type = :question_type,
                                    difficulty = :difficulty,
                                    question_number = :question_number
                                WHERE id = :qid
                            """),
                            {
                                "qid": next_unanswered.id,
                                "question_text": adaptive_question["question_text"],
                                "question_type": adaptive_question["question_type"],
                                "difficulty": adaptive_question["difficulty"],
                                "question_number": next_q_num
                            }
                        )
                    else:
                        connection.execute(
                            text("""
                                INSERT INTO interview_questions (
                                    interview_id, question_number, question_text,
                                    question_type, difficulty
                                )
                                VALUES (
                                    :interview_id, :question_number, :question_text,
                                    :question_type, :difficulty
                                )
                            """),
                            {
                                "interview_id": interview_id,
                                **adaptive_question
                            }
                        )
                    adaptive_added = True

            total_questions = connection.execute(
                text("SELECT COUNT(*) FROM interview_questions WHERE interview_id = :interview_id"),
                {"interview_id": interview_id}
            ).scalar() or (int(answered_count) + (0 if interview_finished else 1))

        return {
            "success": True,
            "answer_id": answer_row.id,
            "wpm": wpm,
            "evaluation": {
                "score": evaluation.get("score"),
                "technical_accuracy": evaluation.get("technical_accuracy"),
                "clarity": evaluation.get("clarity"),
                "problem_solving": evaluation.get("problem_solving"),
                "feedback": evaluation.get("feedback"),
                "what_to_answer": evaluation.get("what_to_answer", ""),
                "recommended_method": evaluation.get("recommended_method", ""),
                "method_guide": evaluation.get("method_guide", ""),
                "strengths": evaluation.get("strengths", []),
                "improvements": evaluation.get("improvements", []),
                "missing_concepts": evaluation.get("missing_concepts", []),
                "interviewer_interest": interviewer_interest,
                "recommend_follow_up": recommend_follow_up,
                "ai_evaluated": evaluation.get("ai_evaluated", False),
                "ai_provider": "gemini",
                "input_mode": input_mode
            },
            "adaptive_added": adaptive_added,
            "finished": interview_finished,
            "answered_count": int(answered_count),
            "total_questions": int(total_questions),
            "message": "Answer evaluated successfully."
        }
    except Exception as error:
        logger.error("Submit answer error: %s", error)
        return {"success": False, "message": "Unable to save your answer."}


def generate_final_report_summary(interview_id):
    try:
        with engine.connect() as connection:
            rows = connection.execute(
                text("""
                    SELECT q.question_text, a.answer_text,
                           a.evaluation_json, a.evaluation_score
                    FROM interview_questions q
                    JOIN interview_answers a ON a.question_id = q.id
                    WHERE q.interview_id = :interview_id
                    ORDER BY q.question_number
                """),
                {"interview_id": interview_id}
            ).fetchall()

        evaluations = []
        for row in rows:
            try:
                parsed = json.loads(row.evaluation_json or "{}")
            except (TypeError, ValueError, json.JSONDecodeError):
                parsed = {}
            if not parsed:
                parsed = {
                    "score": row.evaluation_score,
                    "feedback": "Answer saved and evaluated."
                }
            evaluations.append({
                "question": row.question_text,
                "answer_excerpt": (row.answer_text or "")[:900],
                "evaluation": parsed
            })

        if not evaluations:
            return {
                "strengths": "No answers were submitted for evaluation.",
                "weaknesses": "Complete interview answers to receive meaningful feedback.",
                "suggestions": "Take another interview and answer each question with structured examples."
            }

        prompt = f"""
You are an executive interviewer summarizing the completed interview session.
Evaluations:
{json.dumps(evaluations, ensure_ascii=False)[:14000]}
Return ONLY valid JSON with exactly these string keys: strengths, weaknesses, suggestions.
Be specific, constructive, and concise (1-3 sentences per field).
"""
        result = call_ai_json(prompt)
        if result:
            return {
                "strengths": str(result.get("strengths", "Review the answer-level feedback."))[:2000],
                "weaknesses": str(result.get("weaknesses", "Review the answer-level feedback."))[:2000],
                "suggestions": str(result.get("suggestions", "Practice structured answers with examples."))[:2000]
            }

        scores = [
            float(item["evaluation"].get("score", 0) or 0)
            for item in evaluations
            if item["evaluation"].get("score") is not None
        ]
        avg = round(sum(scores) / len(scores)) if scores else 0
        strength_items = []
        improvement_items = []
        for item in evaluations:
            ev = item["evaluation"]
            strength_items.extend(ev.get("strengths", []) if isinstance(ev.get("strengths"), list) else [])
            improvement_items.extend(ev.get("improvements", []) if isinstance(ev.get("improvements"), list) else [])
        return {
            "strengths": "; ".join(dict.fromkeys(str(x) for x in strength_items if x))[:1800] or "Your answers demonstrated core competence.",
            "weaknesses": "; ".join(dict.fromkeys(str(x) for x in improvement_items if x))[:1800] or "Continue refining system architecture trade-offs.",
            "suggestions": f"Average score was {avg}/100. Review recommended frameworks and model answers."
        }
    except Exception as error:
        logger.error("Final report summary error: %s", error)
        return {
            "strengths": "Your interview answers were saved.",
            "weaknesses": "Summary could not be compiled.",
            "suggestions": "Review question-level feedback."
        }


@app.post("/api/interview/{interview_id}/finish")
def finish_interview(request: Request, interview_id: int):
    user_id = request.session.get("user_id")
    if not user_id:
        return {"success": False, "message": "Please login first.", "redirect": "/login"}

    try:
        with engine.connect() as connection:
            owns_interview = connection.execute(
                text("SELECT id FROM interviews WHERE id = :interview_id AND user_id = :user_id"),
                {"interview_id": interview_id, "user_id": user_id}
            ).fetchone()

        if not owns_interview:
            return {"success": False, "message": "Interview not found."}

        report_summary = generate_final_report_summary(interview_id)

        with engine.begin() as connection:
            interview = connection.execute(
                text("SELECT id FROM interviews WHERE id = :interview_id AND user_id = :user_id"),
                {"interview_id": interview_id, "user_id": user_id}
            ).fetchone()

            if not interview:
                return {"success": False, "message": "Interview not found."}

            stats = connection.execute(
                text("""
                    SELECT
                        COUNT(a.id) AS answered_count,
                        COALESCE(AVG(CASE WHEN LENGTH(TRIM(a.answer_text)) > 0 THEN 1 ELSE 0 END) * 100, 0) AS accuracy_score,
                        COALESCE(AVG(a.words_per_minute), 0) AS average_wpm,
                        AVG(a.evaluation_score) AS evaluation_average
                    FROM interview_questions q
                    LEFT JOIN interview_answers a ON a.question_id = q.id
                    WHERE q.interview_id = :interview_id
                """),
                {"interview_id": interview_id}
            ).fetchone()

            answered_count = int(stats.answered_count or 0)
            accuracy_score = float(stats.accuracy_score or 0)
            speed_score = 100.0

            if stats.average_wpm:
                average_wpm = float(stats.average_wpm)
                speed_score = max(0.0, min(100.0, 100 - abs(120 - average_wpm) * 0.8))

            eval_rows = connection.execute(
                text("""
                    SELECT a.evaluation_score, a.evaluation_json
                    FROM interview_answers a
                    JOIN interview_questions q ON q.id = a.question_id
                    WHERE q.interview_id = :interview_id
                """),
                {"interview_id": interview_id}
            ).fetchall()

            eval_scores = []
            tech_scores = []
            comm_scores = []
            appr_scores = []
            for er in eval_rows:
                if er.evaluation_score is not None:
                    eval_scores.append(float(er.evaluation_score))
                try:
                    ej = json.loads(er.evaluation_json or "{}")
                    if ej.get("technical_accuracy") is not None:
                        tech_scores.append(float(ej["technical_accuracy"]))
                    if ej.get("clarity") is not None:
                        comm_scores.append(float(ej["clarity"]))
                    if ej.get("problem_solving") is not None:
                        appr_scores.append(float(ej["problem_solving"]))
                except Exception:
                    pass

            if eval_scores:
                overall_score = round(sum(eval_scores) / len(eval_scores), 1)
            elif stats.evaluation_average is not None:
                overall_score = round(float(stats.evaluation_average), 1)
            else:
                overall_score = round(accuracy_score, 1)

            final_technical_score = round(sum(tech_scores) / len(tech_scores), 1) if tech_scores else round(accuracy_score, 1)
            final_communication_score = round(sum(comm_scores) / len(comm_scores), 1) if comm_scores else round(speed_score, 1)
            final_approach_score = round(sum(appr_scores) / len(appr_scores), 1) if appr_scores else round(overall_score, 1)
            final_accuracy_score = round(accuracy_score, 1)
            final_speed_score = round(speed_score, 1)

            connection.execute(
                text("""
                    UPDATE interviews
                    SET status = 'completed', completed_at = CURRENT_TIMESTAMP
                    WHERE id = :interview_id
                """),
                {"interview_id": interview_id}
            )

            existing_report = connection.execute(
                text("SELECT id FROM interview_reports WHERE interview_id = :interview_id"),
                {"interview_id": interview_id}
            ).fetchone()

            report_params = {
                "interview_id": interview_id,
                "overall_score": overall_score,
                "technical_score": final_technical_score,
                "communication_score": final_communication_score,
                "accuracy_score": final_accuracy_score,
                "approach_score": final_approach_score,
                "speed_score": final_speed_score,
                "strengths": report_summary["strengths"],
                "weaknesses": report_summary["weaknesses"],
                "suggestions": report_summary["suggestions"]
            }

            if existing_report:
                connection.execute(
                    text("""
                        UPDATE interview_reports
                        SET overall_score = :overall_score,
                            technical_score = :technical_score,
                            communication_score = :communication_score,
                            accuracy_score = :accuracy_score,
                            approach_score = :approach_score,
                            speed_score = :speed_score,
                            strengths = :strengths,
                            weaknesses = :weaknesses,
                            suggestions = :suggestions
                        WHERE interview_id = :interview_id
                    """),
                    report_params
                )
            else:
                connection.execute(
                    text("""
                        INSERT INTO interview_reports (
                            interview_id, overall_score, technical_score,
                            communication_score, accuracy_score, approach_score,
                            speed_score, strengths, weaknesses, suggestions
                        )
                        VALUES (
                            :interview_id, :overall_score, :technical_score,
                            :communication_score, :accuracy_score, :approach_score,
                            :speed_score, :strengths, :weaknesses, :suggestions
                        )
                    """),
                    report_params
                )

        return {
            "success": True,
            "message": "Interview session completed.",
            "interview_id": interview_id,
            "report": {
                "overall_score": overall_score,
                "answered": answered_count,
                "strengths": report_summary["strengths"],
                "weaknesses": report_summary["weaknesses"],
                "suggestions": report_summary["suggestions"]
            }
        }
    except Exception as error:
        logger.error("Finish interview error: %s", error)
        return {"success": False, "message": "Unable to finish the interview."}


@app.post("/start-interview")
async def start_interview(
    request: Request,
    role: str = Form(...),
    custom_role: str = Form(""),
    level: str = Form(...),
    interview_type: str = Form(...),
    resume: UploadFile | None = File(None),
    job_description_file: UploadFile | None = File(None),
    mode: str = Form(...),
    question_count: int | None = Form(None)
):
    user_id = request.session.get("user_id")
    if not user_id:
        return {"success": False, "message": "Please login before starting an interview.", "redirect": "/login"}

    try:
        role = role.strip()
        custom_role = custom_role.strip()
        level = level.strip()
        interview_type = interview_type.strip()
        mode = mode.strip()

        if role == "Other":
            if not custom_role:
                return {"success": False, "message": "Please enter your custom role."}
            final_role = custom_role
        else:
            if not role:
                return {"success": False, "message": "Please select a job role."}
            final_role = role

        if not level:
            return {"success": False, "message": "Please select your experience level."}
        if not interview_type:
            return {"success": False, "message": "Please select an interview type."}
        if not mode:
            return {"success": False, "message": "Please select an interview mode."}
        if resume is None:
            return {"success": False, "message": "Please upload your resume."}

        resume_text = await extract_file_text(resume)
        if not resume_text:
            return {"success": False, "message": "The uploaded resume does not contain readable text."}

        if job_description_file is None:
            return {"success": False, "message": "Please upload a job description."}

        jd_file_text = await extract_file_text(job_description_file)
        if not jd_file_text:
            return {"success": False, "message": "The uploaded job description does not contain readable text."}

        resume_filename = resume.filename if resume.filename else "Resume"
        jd_filename = job_description_file.filename if job_description_file.filename else "Job Description"

        with engine.begin() as connection:
            resume_result = connection.execute(
                text("""
                    INSERT INTO resumes (user_id, filename, extracted_text)
                    VALUES (:user_id, :filename, :extracted_text)
                    RETURNING id
                """),
                {"user_id": user_id, "filename": resume_filename, "extracted_text": resume_text}
            )
            resume_id = resume_result.fetchone().id

            jd_result = connection.execute(
                text("""
                    INSERT INTO job_descriptions (user_id, filename, description_text)
                    VALUES (:user_id, :filename, :description_text)
                    RETURNING id
                """),
                {"user_id": user_id, "filename": jd_filename, "description_text": jd_file_text}
            )
            job_description_id = jd_result.fetchone().id

            interview_result = connection.execute(
                text("""
                    INSERT INTO interviews (user_id, role, experience_level, interview_type, mode, question_count, status)
                    VALUES (:user_id, :role, :experience_level, :interview_type, :mode, 8, 'created')
                    RETURNING id
                """),
                {
                    "user_id": user_id,
                    "role": final_role,
                    "experience_level": level,
                    "interview_type": interview_type,
                    "mode": mode
                }
            )
            interview_id = interview_result.fetchone().id

        question_source = generate_questions_with_fallback(
            interview_id, final_role, level, interview_type, mode, resume_text, jd_file_text, 3
        )

        return {
            "success": True,
            "message": "Interview setup saved successfully.",
            "interview_id": interview_id,
            "resume_id": resume_id,
            "job_description_id": job_description_id,
            "question_source": question_source,
            "interview": {"role": final_role, "level": level, "interview_type": interview_type, "mode": mode},
            "resume": {"filename": resume_filename, "characters": len(resume_text), "text_preview": resume_text[:500]},
            "job_description": {"filename": jd_filename, "characters": len(jd_file_text), "text_preview": jd_file_text[:500]}
        }
    except ValueError as error:
        return {"success": False, "message": str(error)}
    except Exception as error:
        logger.error("Start interview database error: %s", error)
        return {"success": False, "message": "Something went wrong while saving your interview.", "error": str(error)}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=True)